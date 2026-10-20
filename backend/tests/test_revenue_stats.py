import io
import os
import tempfile
import unittest

from openpyxl import load_workbook

from backend.app import create_app, db
from backend.app.models import Drug, Patient, Payment, PrescriptionItem, User, Visit
from backend.app.services.revenue import allocate_payment_revenue, allocate_revenue
from backend.app.services.revenue_query import fetch_revenue_weights
from backend.app.services.time_utils import local_now


class RevenueStatsTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        db_path = os.path.join(self.temp_dir.name, "app.db")

        class TestConfig:
            TESTING = True
            SECRET_KEY = "test-secret-key-that-is-long-enough"
            JWT_SECRET_KEY = "test-jwt-key-that-is-long-enough"
            SQLALCHEMY_DATABASE_URI = "sqlite:///" + db_path
            SQLALCHEMY_TRACK_MODIFICATIONS = False
            CORS_ORIGINS = []

        self.app = create_app(TestConfig)
        self.ctx = self.app.app_context()
        self.ctx.push()
        self.client = self.app.test_client()

        admin = User(username="admin-revenue", real_name="管理员", role="admin")
        admin.set_password("123456")
        self.doctor = User(username="doctor-revenue", real_name="医生", role="doctor")
        self.doctor.set_password("123456")
        self.nurse = User(username="nurse-revenue", real_name="护士", role="nurse")
        self.nurse.set_password("123456")
        db.session.add_all([admin, self.doctor, self.nurse])
        db.session.commit()
        self.admin = admin

        login = self.client.post(
            "/api/auth/login",
            json={"username": "admin-revenue", "password": "123456"},
        )
        self.headers = {"Authorization": f"Bearer {login.get_json()['access_token']}"}

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        db.engine.dispose()
        self.ctx.pop()
        self.temp_dir.cleanup()

    def _create_payment(self, name, amount=10, consult=5, explicit=True):
        patient = Patient(name=name, student_id=name, patient_type="student")
        visit = Visit(
            patient=patient,
            doctor_id=self.doctor.id,
            consultation_fee=consult,
            diagnosis="测试诊断",
            chief_complaint="测试主诉",
        )
        db.session.add(visit)
        db.session.flush()

        definitions = [
            (1, 10, 4),
            (2, 5, 0),
            (3, 5, 2),
        ]
        items = []
        for drug_type, item_amount, cost in definitions:
            drug = Drug(name=f"{name}-{drug_type}", type=drug_type, price=1, status=1)
            item = PrescriptionItem(
                visit_id=visit.id,
                drug=drug,
                amount=item_amount,
                new_amount=item_amount,
                purchase_cost=cost,
            )
            db.session.add(item)
            items.append(item)

        payment = Payment(
            visit_id=visit.id,
            nurse_id=self.nurse.id,
            amount=amount,
            payment_method="cash",
            actual_consultation_fee=2 if explicit else None,
            actual_drug_amount=8 if explicit else None,
        )
        db.session.add(payment)
        db.session.commit()
        return visit, payment, items

    def test_aggregated_weights_match_item_allocation(self):
        visit, payment, items = self._create_payment("张三")

        weights = fetch_revenue_weights([visit.id])[visit.id]
        aggregated = allocate_revenue(
            payment.amount,
            visit.consultation_fee,
            weights,
            weights["cost"],
            explicit_consult=payment.actual_consultation_fee,
            explicit_items=payment.actual_drug_amount,
        )
        expected = allocate_payment_revenue(payment, visit, items)

        self.assertEqual(aggregated, expected)

    def test_revenue_endpoint_uses_aggregated_values_and_returns_page(self):
        self._create_payment("张三")

        response = self.client.get(
            "/api/admin/statistics/revenue",
            headers=self.headers,
            query_string={
                "start_time": local_now().strftime("%Y-%m-%d 00:00:00"),
                "end_time": local_now().strftime("%Y-%m-%d 23:59:59"),
                "page": 1,
                "per_page": 20,
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]

        self.assertEqual(data["total_revenue"], 10.0)
        self.assertEqual(data["consultation_revenue"], 2.0)
        self.assertEqual(data["drug_revenue"], 4.0)
        self.assertEqual(data["service_revenue"], 2.0)
        self.assertEqual(data["consumable_revenue"], 2.0)
        self.assertEqual(data["total_cost"], 6.0)
        self.assertEqual(data["total_profit"], 4.0)
        self.assertEqual(data["total"], 1)
        self.assertEqual(len(data["details"]), 1)
        self.assertEqual(data["details"][0]["drug_amount"], 4.0)

    def test_revenue_endpoint_only_builds_requested_detail_page(self):
        self._create_payment("张三")
        self._create_payment("李四")

        response = self.client.get(
            "/api/admin/statistics/revenue",
            headers=self.headers,
            query_string={
                "start_time": local_now().strftime("%Y-%m-%d 00:00:00"),
                "end_time": local_now().strftime("%Y-%m-%d 23:59:59"),
                "page": 2,
                "per_page": 1,
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]
        self.assertEqual(data["total"], 2)
        self.assertEqual(len(data["details"]), 1)

    def test_revenue_details_are_stably_paginated_newest_first(self):
        _, first_payment, _ = self._create_payment("张三")
        _, second_payment, _ = self._create_payment("李四")

        query_string = {
            "start_time": local_now().strftime("%Y-%m-%d 00:00:00"),
            "end_time": local_now().strftime("%Y-%m-%d 23:59:59"),
            "per_page": 1,
        }
        first_page = self.client.get(
            "/api/admin/statistics/revenue",
            headers=self.headers,
            query_string={**query_string, "page": 1},
        ).get_json()["data"]
        second_page = self.client.get(
            "/api/admin/statistics/revenue",
            headers=self.headers,
            query_string={**query_string, "page": 2},
        ).get_json()["data"]

        self.assertEqual(first_page["total"], 2)
        self.assertEqual(first_page["details"][0]["visit_id"], second_payment.visit_id)
        self.assertEqual(second_page["details"][0]["visit_id"], first_payment.visit_id)

    def test_finance_revenue_by_type_uses_aggregated_values(self):
        self._create_payment("赵六")

        response = self.client.get(
            "/api/finance/revenue/by-type",
            headers=self.headers,
            query_string={
                "start_time": local_now().strftime("%Y-%m-%d 00:00:00"),
                "end_time": local_now().strftime("%Y-%m-%d 23:59:59"),
            },
        )
        self.assertEqual(response.status_code, 200)
        data = response.get_json()["data"]

        self.assertEqual(data["consultation_revenue"], 2.0)
        self.assertEqual(data["drug_revenue"], 4.0)
        self.assertEqual(data["service_revenue"], 2.0)
        self.assertEqual(data["consumable_revenue"], 2.0)
        self.assertEqual(data["total"], 10.0)

    def test_revenue_export_uses_aggregated_values(self):
        self._create_payment("王五")

        response = self.client.get(
            "/api/admin/statistics/revenue/export",
            headers=self.headers,
            query_string={
                "start_time": local_now().strftime("%Y-%m-%d 00:00:00"),
                "end_time": local_now().strftime("%Y-%m-%d 23:59:59"),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        workbook = load_workbook(io.BytesIO(response.data))
        summary = {
            row[0]: row[1]
            for row in workbook["summary"].iter_rows(values_only=True)
            if row and row[0]
        }
        self.assertEqual(summary["总收入"], 10)
        self.assertEqual(summary["总成本"], 6)
        self.assertEqual(summary["总利润"], 4)


if __name__ == "__main__":
    unittest.main()
